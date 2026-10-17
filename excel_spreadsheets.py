import openpyxl as xl
from openpyxl.chart import BarChart, Reference
import pathlib

def process_workbook(filename, output = None, discount = 0.9, sheet_name = 'Sheet1'):
    wb = xl.load_workbook(filename)
    sheet = wb[sheet_name]

    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row, 3)
        # print(cell.value)
        corrected_price = cell.value * discount
        corrected_price_cell = sheet.cell(row, 4)
        corrected_price_cell.value = corrected_price

    values = Reference(sheet, min_row = 2, max_row = sheet.max_row,
                        min_col = 4, max_col = 4)

    chart = BarChart()
    chart.add_data(values)

    sheet.add_chart(chart, 'e2')

    if output is None:
        input_path = pathlib.Path(filename)
        output = input_path.with_name(input_path.stem + '_processed' + input_path.suffix)
    
    # Save to the specified output file
    wb.save(output)

    return output

if __name__ == '__main__':
    print(process_workbook('excel.xlsx'))
    # 生成 excel_processed.xlsx，excel.xlsx 原文件不变

    print(process_workbook('excel.xlsx', output='out.xlsx'))
    # 生成 out.xlsx

    print(process_workbook('excel.xlsx', output='out2.xlsx', discount=0.8))
    # out2.xlsx 里 D 列是 price * 0.8
